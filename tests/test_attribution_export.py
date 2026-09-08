import io
import unittest

from openpyxl import load_workbook

from valuation_app.attribution_export import export_attribution


class AttributionExportTest(unittest.TestCase):
    def payload(self):
        aggregate = {"name": "底层A", "price_impact": 10, "position_amount": 6,
                     "market_change": 16, "formula_difference": 0, "count": 1}
        return {"start": "2026-08-01", "end": "2026-08-31", "scope": "all",
                "summary": {"price_impact": 10, "position_amount": 6, "market_change": 16,
                            "formula_difference": 0, "product_profit": 12, "unattributed": 2,
                            "attributable": 1, "partial": 0, "unavailable": 0},
                "products": [{"name": "FOF", "profit": 12, "price_impact": 10,
                              "unattributed": 2, "opening": 100, "ending": 112}],
                "underlying": [aggregate], "strategies": [dict(aggregate, name="股票指增")],
                "managers": [dict(aggregate, name="管理人")],
                "actions": [dict(aggregate, name="增持")],
                "intervals": [{"fof": "FOF", "name": "底层A", "code": "A", "strategy": "股票指增",
                               "manager": "管理人", "start_date": "2026-08-01", "end_date": "2026-08-02",
                               "action": "增持", "q0": 10, "p0": 2, "q1": 12, "p1": 3,
                               "mv0": 20, "mv1": 36, "price_impact": 10, "position_amount": 6,
                               "market_change": 16, "formula_difference": 0, "quality": "可归属", "note": ""}],
                "rule_note": "估算归因"}

    def test_export_contains_all_evidence_sheets(self):
        book = load_workbook(io.BytesIO(export_attribution(self.payload())), read_only=True)
        self.assertEqual(book.sheetnames, ["归因概览", "产品勾稽", "底层汇总", "策略汇总", "管理人汇总", "动作汇总", "逐区间明细"])
        self.assertEqual(book["逐区间明细"].max_row, 2)

    def test_export_rejects_unbounded_detail(self):
        value = self.payload()
        value["intervals"] = [{}] * 20001
        with self.assertRaises(ValueError):
            export_attribution(value)


if __name__ == "__main__":
    unittest.main()
