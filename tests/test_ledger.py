import unittest
import os
import tempfile
from datetime import date, datetime

import openpyxl

from valuation_app.ledger import load_ledger, load_product_metadata, parse_date
from valuation_app.parser import scan_valuations
from valuation_app.config import match_product


class LedgerTest(unittest.TestCase):
    def test_parse_mixed_dates(self):
        self.assertEqual(parse_date(datetime(2026, 7, 1)), date(2026, 7, 1))
        self.assertEqual(parse_date(20240613), date(2024, 6, 13))
        self.assertEqual(parse_date("2026.1.7"), date(2026, 1, 7))
        self.assertIsNone(parse_date("-"))

    def test_ledger_short_aliases_match_continuation_rows(self):
        expected = {
            "天玑13": "第一创业天玑13号单一资产管理计划",
            "盛乾同行2": "国金资管盛乾同行2号FOF单一资产管理计划",
            "源泉优享FOF3": "第一创业源泉优享FOF3号单一资产管理计划",
        }
        for alias, product in expected.items():
            self.assertEqual(match_product(alias), product)

    def test_blank_ledger_continuation_row_inherits_product(self):
        workbook = openpyxl.Workbook()
        workbook.remove(workbook.active)
        for year in ("2024", "2025", "2026"):
            sheet = workbook.create_sheet(year)
            sheet.cell(1, 3).value = "项目名称"
            sheet.cell(1, 6).value = "投资日期"
            sheet.cell(1, 7).value = "投资金额"
        sheet = workbook["2024"]
        sheet.cell(2, 3).value = "第一创业天玑13号单一资产管理计划入池并投资"
        sheet.cell(2, 6).value = 20240613
        sheet.cell(2, 7).value = 1000
        sheet.cell(3, 6).value = 20240704
        sheet.cell(3, 7).value = 2000
        handle, path = tempfile.mkstemp(suffix=".xlsx")
        os.close(handle)
        try:
            workbook.save(path)
            flows, errors = load_ledger(path)
            self.assertEqual(errors, [])
            self.assertEqual(len(flows), 2)
            self.assertEqual(sum(flow.amount for flow in flows), 30000000)
            self.assertEqual(flows[0].product, flows[1].product)
        finally:
            workbook.close()
            os.remove(path)

    def test_july_known_cashflows(self):
        if not os.path.exists("专户资金台账.xlsx"):
            self.skipTest("workspace has no ledger")
        snapshots, _ = scan_valuations("products")
        flows, errors = load_ledger("专户资金台账.xlsx", snapshots)
        self.assertEqual(errors, [])
        expected = {
            "第一创业源泉优享FOF3号单一资产管理计划": 10000000,
            "国金资管盛乾同行2号FOF单一资产管理计划": 22000000,
            "第一创业天玑13号单一资产管理计划": -10000000,
        }
        for product, amount in expected.items():
            actual = sum(flow.amount for flow in flows if flow.product == product and
                         "2026-07-01" <= flow.flow_date <= "2026-07-31")
            self.assertEqual(actual, amount)

    def test_2026_product_metadata(self):
        if not os.path.exists("专户资金台账.xlsx"):
            self.skipTest("workspace has no ledger")
        metadata, errors = load_product_metadata("专户资金台账.xlsx")
        self.assertEqual(errors, [])
        self.assertEqual(metadata["第一创业天玑13号单一资产管理计划"]["total_investment"], 806000000)
        self.assertEqual(metadata["第一创业天玑13号单一资产管理计划"]["approved_quota"], 300000000)
        self.assertEqual(metadata["西南证券嘉盈1号FOF单一资产管理计划"]["approved_quota"], 3000000000)


if __name__ == "__main__":
    unittest.main()
