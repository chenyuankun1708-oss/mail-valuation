import os
import tempfile
import unittest

from openpyxl import Workbook

from valuation_app.otc_ledger import migrate_workbook
from valuation_app.knowledge import MAX_FILE_SIZE
from valuation_app.otc_store import (ATTACHMENT_CATEGORIES, OtcStore,
                                     RevisionConflictError, TaskConflictError)


class OtcLedgerTest(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.store = OtcStore(self.temporary.name)

    def tearDown(self):
        self.temporary.cleanup()

    def test_product_revision_manual_event_soft_void_and_backtest_mapping(self):
        product = self.store.create_product({
            "name": "测试发行", "strategy_name": "DCN", "structure": "dcn",
            "index_code": "000852", "notional": 10000000, "start_date": "2026-01-02",
            "terms": {"term_months": 24, "lock_period_months": 3,
                      "knock_in_ratio": .7, "knock_out_initial": 1,
                      "knock_out_decrease_monthly": .005, "first_coupon": .12,
                      "second_coupon": .12, "coupon_switch_months": 12,
                      "dividend_barrier": .8, "monthly_dividend": .008},
            "reference": {}, "notes": "仅测试"}, "alice")
        updated = self.store.update_product(product["id"], {"notes": "修订"}, 1, "bob")
        with self.assertRaises(RevisionConflictError):
            self.store.update_product(product["id"], {"notes": "冲突"}, 1, "alice")
        ki = self.store.add_event(product["id"], {"event_type": "敲入确认",
            "event_date": "2026-02-02", "values": {"note": "人工确认"},
            "expected_revision": updated["revision"]}, "bob")
        self.assertEqual(ki["status"], "草稿")
        active = self.store.add_event(product["id"], {"event_type": "生效",
            "event_date": "2026-01-02", "values": {}, "expected_revision": ki["revision"]}, "alice")
        self.assertEqual(active["status"], "存续")
        request = self.store.product_backtest_request(product["id"], active["revision"],
                                                       {"start_date": "2020-01-01", "end_date": "2026-01-01"})
        self.assertEqual(request["product_id"], product["id"])
        self.assertEqual(request["product_name"], "测试发行")
        voided = self.store.void_product(product["id"], active["revision"], "alice")
        self.assertEqual(voided["status"], "已作废")
        self.assertGreaterEqual(len(self.store.get_product(product["id"])["audit"]), 5)

    def test_settled_event_sets_end_date_and_uses_settlement_return_boundary(self):
        product = self.store.create_product({
            "name": "了结收益产品", "strategy_name": "雪球", "structure": "classic_snowball",
            "index_code": "000852", "notional": 1000, "start_date": "2024-02-29",
            "valuation_amount": 1120, "cumulative_distributions": 20,
            "performance_hurdle": .05, "performance_fee_rate": .20,
            "terms": {"term_months": 24}, "reference": {}, "notes": ""}, "alice", status="存续")
        settled = self.store.add_event(product["id"], {
            "event_type": "敲出确认", "event_date": "2025-02-28", "values": {},
            "expected_revision": product["revision"]}, "alice")
        self.assertEqual(settled["status"], "已敲出")
        self.assertEqual(settled["end_date"], "2025-02-28")
        self.assertEqual(settled["returns"]["basis"], "settlement")
        self.assertEqual(settled["returns"]["actual_end_date"], "2025-02-28")
        self.assertAlmostEqual(settled["returns"]["gross_absolute_return"], .14)
        self.assertAlmostEqual(settled["returns"]["gross_annualized_return"], .14)

    def test_product_attachments_deduplicate_extract_and_soft_delete(self):
        self.assertEqual(set(ATTACHMENT_CATEGORIES), {
            "产品合同", "投资者名单", "公告", "回测报告", "宣传材料", "其他"})
        first = self.store.create_product({"name": "附件产品A", "terms": {}, "reference": {}}, "alice")
        second = self.store.create_product({"name": "附件产品B", "terms": {}, "reference": {}}, "alice")
        source = os.path.join(self.temporary.name, "contract.md")
        with open(source, "w", encoding="utf-8") as output:
            output.write("名义本金：1000万元")
        attachment, duplicate = self.store.add_attachment(
            first["id"], source, "合同.md", "产品合同", "alice")
        self.assertFalse(duplicate)
        same, duplicate = self.store.add_attachment(
            first["id"], source, "合同副本.md", "产品合同", "alice")
        self.assertTrue(duplicate)
        self.assertEqual(same["id"], attachment["id"])
        other, duplicate = self.store.add_attachment(
            second["id"], source, "合同.md", "产品合同", "bob")
        self.assertFalse(duplicate)
        self.assertNotEqual(other["id"], attachment["id"])
        self.assertEqual(other["sha256"], attachment["sha256"])
        detail = self.store.get_attachment(first["id"], attachment["id"], True)
        self.assertIn("名义本金", detail["text_content"])
        self.assertTrue(os.path.isfile(self.store.attachment_file_path(
            first["id"], attachment["id"])[0]))
        investor, _ = self.store.add_attachment(
            first["id"], source, "名单.md", "投资者名单", "alice")
        _, context = self.store.field_audit_context(first["id"], first["revision"])
        self.assertEqual([item["id"] for item in context], [attachment["id"]])
        self.assertNotIn(investor["id"], [item["id"] for item in context])
        self.store.deactivate_attachment(first["id"], attachment["id"], 1, "alice")
        self.assertEqual(self.store.list_attachments(first["id"])[0]["id"], investor["id"])
        with self.assertRaises(KeyError):
            self.store.get_attachment(first["id"], attachment["id"])
        oversized = os.path.join(self.temporary.name, "oversized.txt")
        with open(oversized, "wb") as output:
            output.seek(MAX_FILE_SIZE)
            output.write(b"x")
        with self.assertRaises(ValueError):
            self.store.add_attachment(first["id"], oversized, "oversized.txt", "其他")

    def test_latest_completed_field_audit_is_returned_with_product_detail(self):
        product = self.store.create_product({
            "name": "排查产品", "strategy_name": "DCN", "terms": {},
            "reference": {}, "notes": ""}, "alice")
        audit = self.store.create_field_audit(product["id"], "assistant")
        result = {"fields": [{"field": "期限", "status": "不一致",
                              "input_value": "24个月", "document_value": "36个月",
                              "evidence": [{"location": "宣传材料", "quote": "合同期限36个月"}]}]}
        self.store.finish_field_audit(
            audit["id"], "completed", "assistant_manual", "codex", result)
        detail = self.store.get_product(product["id"])
        self.assertEqual(detail["latest_field_audit"]["id"], audit["id"])
        self.assertEqual(detail["latest_field_audit"]["result"], result)
        self.assertEqual(detail["latest_field_audit"]["provider"], "assistant_manual")

    def test_performance_fee_returns_filters_aggregates_and_templates(self):
        product = self.store.create_product({
            "name": "收益产品", "strategy_name": "雪球", "structure": "classic_snowball",
            "index_code": "000852", "notional": 1000, "start_date": "2024-01-01",
            "return_as_of_date": "2025-01-01", "valuation_amount": 1150,
            "cumulative_distributions": 10, "performance_hurdle": .05,
            "performance_fee_rate": .20, "terms": {"term_months": 24},
            "reference": {}, "notes": ""}, "alice", status="存续")
        returns = product["returns"]
        self.assertAlmostEqual(returns["gross_profit"], 160)
        self.assertAlmostEqual(returns["performance_fee"], (160 - 1000 * .05 * 366 / 365) * .2)
        self.assertEqual(returns["status"], "complete")
        listing = self.store.list_products(1, 20, {"lifecycle": "存续", "return_state": "complete"})
        self.assertEqual(listing["total"], 1)
        self.assertAlmostEqual(listing["aggregates"]["sums"]["gross_profit"], 160)
        template = self.store.create_template("标准雪球", {
            "structure": "classic_snowball", "index_code": "000852",
            "start_date": "2020-01-01", "end_date": "2026-01-01"}, "alice")
        updated = self.store.update_template(template["id"], "标准雪球2", template["request"], 1, "bob")
        self.assertEqual(updated["revision"], 2)
        self.assertEqual(self.store.void_template(template["id"], 2, "bob")["status"], "void")

    def test_migration_keeps_13_rows_duplicates_raw_trace_and_is_idempotent(self):
        source = os.path.join(self.temporary.name, "历史零售雪球要素.xlsx")
        workbook = Workbook(); sheet = workbook.active; sheet.title = "Sheet1"
        for row in range(3, 16):
            sheet.cell(row, 1, "重复产品" if row in (9, 10) else "产品%d" % row)
            sheet.cell(row, 2, "原始策略%d" % row); sheet.cell(row, 3, "000852.SH")
            sheet.cell(row, 4, 1000 + row); sheet.cell(row, 5, 24)
            sheet.cell(row, 6, "2025-01-%02d" % row); sheet.cell(row, 7, 6000 + row)
            sheet.cell(row, 15, "存续" if row < 7 else "2026-01-01")
            sheet.cell(row, 17, .08 if row >= 7 else None)
            sheet.cell(row, 18, "原始备注")
        workbook.save(source)
        with open(source, "rb") as handle:
            before = handle.read()
        result = migrate_workbook(self.store, source)
        self.assertEqual(result["imported"], 13)
        with open(source, "rb") as handle:
            self.assertEqual(handle.read(), before)
        products = self.store.list_products(1, 50)["items"]
        self.assertEqual(sum(item["name"] == "重复产品" for item in products), 2)
        self.assertEqual(sum(item["status"] == "存续" for item in products), 4)
        self.assertTrue(all(item["structure"] is None for item in products))
        traced = next(item for item in products if item["source_row"] == 3)
        self.assertEqual(traced["reference"]["source_row"], 3)
        ended = next(item for item in products if item["source_row"] == 7)
        self.assertEqual(ended["end_date"], "2026-01-01")
        self.assertEqual(ended["legacy_annual_return"], .08)
        self.assertIsNone(traced["end_date"])
        with self.assertRaises(TaskConflictError):
            migrate_workbook(self.store, source)

    def test_wide_ledger_filters_sorting_matching_ids_and_selected_aggregate(self):
        first = self.store.create_product({
            "name": "B产品", "strategy_name": "原始DCN", "structure": "dcn",
            "index_code": "000852", "notional": 1000, "start_date": "2025-01-01",
            "end_date": "2026-01-02", "return_as_of_date": "2026-01-02",
            "valuation_amount": 1100, "legacy_annual_return": .12,
            "performance_hurdle": 0, "performance_fee_rate": 0,
            "terms": {"term_months": 12}, "reference": {}, "notes": ""}, "alice")
        second = self.store.create_product({
            "name": "A产品", "strategy_name": "原始雪球", "structure": "classic_snowball",
            "index_code": "000905", "notional": 2000, "start_date": "2024-01-01",
            "end_date": None, "return_as_of_date": None, "valuation_amount": None,
            "legacy_annual_return": None, "terms": {"term_months": 24},
            "reference": {}, "notes": ""}, "alice")
        listing = self.store.list_products(1, 20, {
            "strategy_name": "原始DCN", "end_from": "2026-01-01",
            "end_to": "2026-12-31", "sort_by": "gross_profit", "sort_dir": "desc"})
        self.assertEqual(listing["matching_ids"], [first["id"]])
        self.assertEqual(listing["items"][0]["legacy_annual_return"], .12)
        selected = self.store.aggregate_products([first["id"]])
        self.assertEqual(selected["count"], 1)
        self.assertAlmostEqual(selected["sums"]["gross_profit"], 100)
        self.assertEqual(self.store.list_products(1, 20, {
            "sort_by": "name", "sort_dir": "asc"})["items"][0]["id"], second["id"])
        with self.assertRaises(ValueError):
            self.store.aggregate_products(["../secret"])
