"""Restricted OTC option backtests using the local Wind index cache.

The calculation layer is intentionally independent from the web server.  It
contains no database credentials, network access, simulated-price fallback or
plotting dependency.  Calendar observations follow the source project's
calendar_365 implementation at commit 1df9e0b plus working-tree patch
59dfa1e66c54607d1d5e6f9228e6de353d8084d0.
"""

from __future__ import division

import hashlib
import json
import math
import os
from dataclasses import dataclass
from datetime import datetime

import numpy as np
import pandas as pd
from openpyxl import Workbook


ENGINE_VERSION = "otc-calendar-365-v2"
REPORT_SCHEMA_VERSION = 2
SOURCE_COMMIT = "1df9e0be3f1da75221172c71f60905ed7c247e3d"
SOURCE_PATCH_HASH = "59dfa1e66c54607d1d5e6f9228e6de353d8084d0"
STRUCTURES = {
    "classic_snowball": "经典雪球",
    "european_snowball": "欧式雪球",
    "dcn": "DCN",
    "dcn_snowball_combo": "DCN＋雪球组合",
}
INDICES = {
    "000852": "中证1000",
    "000905": "中证500",
    "000300": "沪深300",
}

PARAMETER_DEFINITIONS = (
    {"name": "term_months", "label": "期限（月）", "type": "integer", "min": 1, "max": 60, "default": 24, "structures": "all"},
    {"name": "lock_period_months", "label": "敲出锁定期（月）", "type": "integer", "min": 1, "max": 60, "default": 3, "structures": "all"},
    {"name": "knock_in_ratio", "label": "敲入比例", "type": "number", "min": .30, "max": .99, "step": .001, "default": .70, "structures": "all"},
    {"name": "knock_out_initial", "label": "初始敲出比例", "type": "number", "min": .80, "max": 1.20, "step": .001, "default": 1.0, "structures": "all"},
    {"name": "knock_out_decrease_monthly", "label": "每月降敲", "type": "number", "min": 0, "max": .05, "step": .0001, "default": .005, "structures": "all"},
    {"name": "first_coupon", "label": "前段年化票息", "type": "number", "min": 0, "max": .50, "step": .001, "default": .12, "structures": ("classic_snowball", "european_snowball", "dcn_snowball_combo")},
    {"name": "second_coupon", "label": "后段年化票息", "type": "number", "min": 0, "max": .50, "step": .001, "default": .12, "structures": ("classic_snowball", "european_snowball", "dcn_snowball_combo")},
    {"name": "coupon_switch_months", "label": "票息切换月", "type": "integer", "min": 1, "max": 60, "default": 12, "structures": ("classic_snowball", "european_snowball", "dcn_snowball_combo")},
    {"name": "max_loss", "label": "最大亏损", "type": "number", "min": .01, "max": 1, "step": .01, "default": None, "optional": True, "structures": ("classic_snowball", "european_snowball", "dcn")},
    {"name": "dividend_barrier", "label": "派息障碍", "type": "number", "min": .50, "max": .99, "step": .001, "default": .80, "structures": ("dcn", "dcn_snowball_combo")},
    {"name": "monthly_dividend", "label": "月派息率", "type": "number", "min": 0, "max": .05, "step": .0001, "default": .0088, "structures": ("dcn", "dcn_snowball_combo")},
    {"name": "dcn_weight", "label": "DCN收益组合系数", "type": "number", "min": .01, "max": 10, "step": .01, "default": 1.0, "structures": ("dcn_snowball_combo",)},
    {"name": "snowball_weight", "label": "雪球收益组合系数", "type": "number", "min": .01, "max": 10, "step": .01, "default": .2, "structures": ("dcn_snowball_combo",)},
    {"name": "dcn_max_loss", "label": "DCN最大亏损", "type": "number", "min": .01, "max": 1, "step": .01, "default": None, "optional": True, "structures": ("dcn_snowball_combo",)},
    {"name": "snowball_max_loss", "label": "雪球最大亏损", "type": "number", "min": .01, "max": 1, "step": .01, "default": None, "optional": True, "structures": ("dcn_snowball_combo",)},
)


def parameter_definitions():
    return [dict(item) for item in PARAMETER_DEFINITIONS]


def validate_partial_terms(raw):
    """Validate stored product terms without inventing missing contract terms."""
    if raw in (None, ""):
        return {}
    if not isinstance(raw, dict):
        raise ValueError("terms必须是对象")
    definitions = {item["name"]: item for item in PARAMETER_DEFINITIONS}
    if set(raw) - set(definitions):
        raise ValueError("包含未登记的回测参数")
    clean = {}
    for name, value in raw.items():
        definition = definitions[name]
        if value in (None, "") and definition.get("optional"):
            clean[name] = None
        elif definition["type"] == "integer":
            clean[name] = _integer(value, definition["label"], definition["min"], definition["max"])
        else:
            clean[name] = _number(value, definition["label"], definition["min"], definition["max"])
    return clean


def _number(value, name, minimum=None, maximum=None, optional=False):
    if optional and (value is None or value == ""):
        return None
    if isinstance(value, bool):
        raise ValueError("%s必须是数字" % name)
    try:
        result = float(value)
    except (TypeError, ValueError):
        raise ValueError("%s必须是数字" % name)
    if not math.isfinite(result):
        raise ValueError("%s必须是有限数字" % name)
    if minimum is not None and result < minimum:
        raise ValueError("%s不能小于%s" % (name, minimum))
    if maximum is not None and result > maximum:
        raise ValueError("%s不能大于%s" % (name, maximum))
    return result


def _integer(value, name, minimum, maximum):
    result = _number(value, name, minimum, maximum)
    if result != int(result):
        raise ValueError("%s必须是整数" % name)
    return int(result)


def _iso_date(value, name):
    try:
        parsed = datetime.strptime(str(value), "%Y-%m-%d")
    except (TypeError, ValueError):
        raise ValueError("%s必须是ISO日期" % name)
    return parsed.strftime("%Y-%m-%d")


def validate_request(raw):
    """Return a canonical, bounded request suitable for persistent execution."""
    if not isinstance(raw, dict):
        raise ValueError("请求必须是JSON对象")
    allowed = {
        "structure", "index_code", "start_date", "end_date", "term_months",
        "lock_period_months", "knock_in_ratio", "knock_out_initial",
        "knock_out_decrease_monthly", "first_coupon", "second_coupon",
        "coupon_switch_months", "max_loss", "dividend_barrier",
        "monthly_dividend", "dcn_weight", "snowball_weight", "dcn_max_loss",
        "snowball_max_loss", "product_id", "product_revision", "template_id",
        "template_revision", "product_name",
    }
    extra = set(raw) - allowed
    if extra:
        raise ValueError("包含未允许参数：%s" % ", ".join(sorted(extra)))
    structure = str(raw.get("structure") or "")
    if structure not in STRUCTURES:
        raise ValueError("期权结构不在固定清单")
    index_code = str(raw.get("index_code") or "")
    if index_code not in INDICES:
        raise ValueError("标的指数不在固定清单")
    start_date = _iso_date(raw.get("start_date"), "start_date")
    end_date = _iso_date(raw.get("end_date"), "end_date")
    if start_date >= end_date:
        raise ValueError("开始日期必须早于评价截止日")
    term_months = _integer(raw.get("term_months", 24), "期限（月）", 1, 60)
    lock_months = _integer(raw.get("lock_period_months", 3), "敲出锁定期", 1, term_months)
    switch_months = _integer(raw.get("coupon_switch_months", min(12, term_months)),
                             "票息切换月", 1, term_months)
    result = {
        "structure": structure,
        "index_code": index_code,
        "start_date": start_date,
        "end_date": end_date,
        "term_months": term_months,
        "lock_period_months": lock_months,
        "knock_in_ratio": _number(raw.get("knock_in_ratio", 0.70), "敲入比例", 0.30, 0.99),
        "knock_out_initial": _number(raw.get("knock_out_initial", 1.0), "初始敲出比例", 0.80, 1.20),
        "knock_out_decrease_monthly": _number(raw.get("knock_out_decrease_monthly", 0.005), "每月降敲", 0, 0.05),
        "first_coupon": _number(raw.get("first_coupon", 0.12), "前段年化票息", 0, 0.50),
        "second_coupon": _number(raw.get("second_coupon", 0.12), "后段年化票息", 0, 0.50),
        "coupon_switch_months": switch_months,
        "max_loss": _number(raw.get("max_loss"), "最大亏损", 0.01, 1.0, True),
        "dividend_barrier": _number(raw.get("dividend_barrier", 0.80), "派息障碍", 0.50, 0.99),
        "monthly_dividend": _number(raw.get("monthly_dividend", 0.0088), "月派息率", 0, 0.05),
        "dcn_weight": _number(raw.get("dcn_weight", 1.0), "DCN收益组合系数", 0.01, 10),
        "snowball_weight": _number(raw.get("snowball_weight", 0.2), "雪球收益组合系数", 0.01, 10),
    }
    product_name = str(raw.get("product_name") or "").strip()
    if len(product_name) > 120:
        raise ValueError("产品名称不能超过120字")
    if any(ord(char) < 32 for char in product_name) or any(char in product_name for char in '/\\:*?"<>|'):
        raise ValueError("产品名称不得包含控制字符或路径/文件名特殊字符")
    result["product_name"] = product_name or "%s－%s－%s" % (
        STRUCTURES[structure], INDICES[index_code], end_date)
    if structure == "dcn_snowball_combo":
        result["dcn_max_loss"] = _number(raw.get("dcn_max_loss"), "DCN最大亏损", .01, 1, True)
        result["snowball_max_loss"] = _number(raw.get("snowball_max_loss"), "雪球最大亏损", .01, 1, True)
    product_id = raw.get("product_id")
    if product_id is not None:
        text = str(product_id)
        if len(text) != 36:
            raise ValueError("产品ID无效")
        result["product_id"] = text
        result["product_revision"] = _integer(raw.get("product_revision"), "产品修订号", 1, 1000000000)
    template_id = raw.get("template_id")
    if template_id is not None:
        text = str(template_id)
        if len(text) != 36:
            raise ValueError("模板ID无效")
        result["template_id"] = text
        result["template_revision"] = _integer(raw.get("template_revision"), "模板修订号", 1, 1000000000)
    return result


def _sha256(path):
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def load_prices(cache_path, request):
    """Read only a registered index from the existing local Wind cache."""
    with open(cache_path, "r", encoding="utf-8") as handle:
        payload = json.load(handle)
    if payload.get("source") != "Wind Oracle数据库":
        raise ValueError("行情缓存不是已登记的Wind Oracle数据")
    item = (payload.get("indices") or {}).get(request["index_code"]) or {}
    if not item.get("points"):
        raise ValueError("所选指数没有本地Wind缓存")
    points = [point for point in item["points"]
              if request["start_date"] <= str(point.get("date") or "") <= request["end_date"]]
    rows = []
    for point in points:
        try:
            close = float(point["close"])
            day = datetime.strptime(point["date"], "%Y-%m-%d")
        except (KeyError, TypeError, ValueError):
            continue
        if math.isfinite(close) and close > 0:
            rows.append((day, close))
    if len(rows) < 2:
        raise ValueError("请求区间内少于两个有效行情日")
    frame = pd.DataFrame(rows, columns=["date", "price"]).drop_duplicates("date", keep="last")
    frame = frame.sort_values("date").reset_index(drop=True)
    return frame, {
        "source": payload.get("source"),
        "cache_updated_at": payload.get("updated_at"),
        "index_updated_at": item.get("updated_at"),
        "cache_sha256": _sha256(cache_path),
        "actual_start_date": frame["date"].iloc[0].strftime("%Y-%m-%d"),
        "actual_end_date": frame["date"].iloc[-1].strftime("%Y-%m-%d"),
        "price_rows": len(frame),
    }


def _following(dates, contract_date):
    position = dates.searchsorted(pd.Timestamp(contract_date).to_datetime64(), side="left")
    return None if position >= len(dates) else pd.Timestamp(dates[position])


@dataclass
class Terms:
    term_months: int
    lock_period_months: int
    knock_in_ratio: float
    knock_out_initial: float
    knock_out_decrease_monthly: float
    first_coupon: float
    second_coupon: float
    coupon_switch_months: int
    max_loss: object
    dividend_barrier: float
    monthly_dividend: float


def _schedule(frame, entry_position, terms, positions):
    entry = frame["date"].iloc[entry_position]
    dates = frame["date"].to_numpy()
    maturity_contract = entry + pd.DateOffset(months=terms.term_months)
    maturity = _following(dates, maturity_contract)
    observations = [
        _following(dates, entry + pd.DateOffset(months=month))
        for month in range(terms.lock_period_months, terms.term_months + 1)
    ]
    dividends = [
        _following(dates, entry + pd.DateOffset(months=month))
        for month in range(1, terms.term_months + 1)
    ]
    switch = _following(dates, entry + pd.DateOffset(months=terms.coupon_switch_months))
    return entry, maturity_contract, maturity, observations, dividends, switch


def _base_row(frame, entry_position, exit_position, exit_date, knocked_out, ongoing):
    entry = frame["date"].iloc[entry_position]
    return {
        "entry_date": entry.strftime("%Y-%m-%d"),
        "entry_price": float(frame["price"].iloc[entry_position]),
        "exit_date": exit_date.strftime("%Y-%m-%d") if not ongoing else None,
        "exit_price": float(frame["price"].iloc[exit_position]) if not ongoing else None,
        "holding_calendar_days": int((exit_date - entry).days) if not ongoing else None,
        "holding_trading_days": int(exit_position - entry_position) if not ongoing else None,
        "knocked_out": bool(knocked_out),
        "knocked_in": False,
        "still_running": bool(ongoing),
        "absolute_return": None,
        "annualized_return": None,
        "exit_observation_month": None,
    }


def _single(frame, entry_position, terms, kind):
    positions = {pd.Timestamp(day): pos for pos, day in enumerate(frame["date"])}
    entry, maturity_contract, maturity, observations, dividends, switch = _schedule(
        frame, entry_position, terms, positions)
    entry_price = float(frame["price"].iloc[entry_position])
    ko_date = None
    ko_month = None
    for offset, observation in enumerate(observations):
        if observation is None or observation not in positions:
            continue
        ratio = float(frame["price"].iloc[positions[observation]]) / entry_price
        barrier = terms.knock_out_initial - offset * terms.knock_out_decrease_monthly
        # Preserve the source engine's strict-above treatment of barrier equality.
        if ratio > barrier and not np.isclose(ratio, barrier):
            ko_date = observation
            ko_month = terms.lock_period_months + offset
            break
    if maturity is None and ko_date is None:
        last_position = len(frame) - 1
        return _base_row(frame, entry_position, last_position,
                         frame["date"].iloc[last_position], False, True)
    exit_date = ko_date or maturity
    exit_position = positions[exit_date]
    row = _base_row(frame, entry_position, exit_position, exit_date, ko_date is not None, False)
    row["maturity_contract_date"] = maturity_contract.strftime("%Y-%m-%d")
    row["adjusted_maturity_date"] = maturity.strftime("%Y-%m-%d") if maturity is not None else None
    row["exit_observation_month"] = ko_month if ko_date is not None else terms.term_months
    path = frame["price"].iloc[entry_position:exit_position + 1] / entry_price
    if kind == "european_snowball":
        knocked_in = bool(path.iloc[-1] <= terms.knock_in_ratio)
    else:
        knocked_in = bool(path.min() <= terms.knock_in_ratio)
    row["knocked_in"] = knocked_in
    if kind == "dcn":
        dividend_count = sum(
            1 for day in dividends
            if day is not None and day <= exit_date and
            float(frame["price"].iloc[positions[day]]) / entry_price >= terms.dividend_barrier
        )
        absolute = dividend_count * terms.monthly_dividend
        if ko_date is None and knocked_in:
            absolute += float(path.iloc[-1] - 1)
        row["dividend_count"] = dividend_count
    else:
        if ko_date is None and knocked_in:
            absolute = float(path.iloc[-1] - 1)
        else:
            coupon = terms.first_coupon if switch is None or exit_date <= switch else terms.second_coupon
            absolute = coupon * row["holding_calendar_days"] / 365.0
    if terms.max_loss is not None:
        absolute = max(absolute, -terms.max_loss)
    row["absolute_return"] = float(absolute)
    row["annualized_return"] = float(absolute * 365.0 / max(row["holding_calendar_days"], 1))
    return row


def _terms(request, max_loss=None):
    return Terms(
        request["term_months"], request["lock_period_months"], request["knock_in_ratio"],
        request["knock_out_initial"], request["knock_out_decrease_monthly"],
        request["first_coupon"], request["second_coupon"], request["coupon_switch_months"],
        request["max_loss"] if max_loss is None else max_loss,
        request["dividend_barrier"], request["monthly_dividend"],
    )


def _summary(samples, structure):
    completed = [row for row in samples if not row["still_running"]]
    count = len(completed)
    def average(key):
        values = [row[key] for row in completed if row.get(key) is not None]
        return float(sum(values) / len(values)) if values else None
    summary = {
        "total_samples": len(samples),
        "completed_samples": count,
        "still_running_samples": len(samples) - count,
        "knocked_out_no_knock_in": sum(row["knocked_out"] and not row["knocked_in"] for row in completed),
        "not_knocked_out_no_knock_in": sum(not row["knocked_out"] and not row["knocked_in"] for row in completed),
        "knocked_in_then_out": sum(row["knocked_out"] and row["knocked_in"] for row in completed),
        "not_knocked_out_knocked_in": sum(not row["knocked_out"] and row["knocked_in"] for row in completed),
        "positive_return_probability": (sum(row["absolute_return"] > 0 for row in completed) / count if count else None),
        "first_year_knock_out_probability": (sum(row["knocked_out"] and row["holding_calendar_days"] <= 365 for row in completed) / count if count else None),
        "average_holding_calendar_days": average("holding_calendar_days"),
        "average_absolute_return": average("absolute_return"),
        "average_annualized_return": average("annualized_return"),
    }
    if structure in ("dcn", "dcn_snowball_combo"):
        summary["average_dividend_count"] = average("dividend_count")
    return summary


def _charts(samples):
    completed = [row for row in samples if not row["still_running"]]
    if len(completed) > 500:
        positions = np.linspace(0, len(completed) - 1, 500).astype(int)
        line_rows = [completed[position] for position in positions]
    else:
        line_rows = completed
    histogram = {}
    for row in completed:
        month = int(row.get("exit_observation_month") or 0)
        histogram[str(month)] = histogram.get(str(month), 0) + 1
    return {
        "sample_returns": [{"date": row["entry_date"], "value": row["absolute_return"]}
                           for row in line_rows],
        "holding_months": [{"month": int(key), "count": value}
                           for key, value in sorted(histogram.items(), key=lambda item: int(item[0]))],
    }


def _component_samples(samples, component):
    """Project combo rows into component rows without recalculating a payoff."""
    rows = []
    for source in samples:
        row = dict(source)
        if component == "dcn":
            row["absolute_return"] = source.get("dcn_absolute_return")
            row["knocked_in"] = bool(source.get("dcn_knocked_in"))
        else:
            row["absolute_return"] = source.get("snowball_absolute_return")
            row["knocked_in"] = bool(source.get("snowball_knocked_in"))
            row.pop("dividend_count", None)
        if row.get("absolute_return") is not None and row.get("holding_calendar_days"):
            row["annualized_return"] = row["absolute_return"] * 365.0 / row["holding_calendar_days"]
        else:
            row["annualized_return"] = None
        rows.append(row)
    return rows


def _report_pages(structure):
    if structure == "dcn_snowball_combo":
        return [
            "组合情景路径", "DCN情景路径", "雪球情景路径",
            "DCN参数", "DCN汇总", "DCN收益图", "DCN存续期", "DCN派息分布",
            "雪球参数", "雪球汇总", "雪球收益图", "雪球存续期",
            "组合参数", "组合汇总", "组合收益图", "组合存续期",
        ]
    pages = ["条款情景路径", "参数", "汇总", "指数与样本收益", "存续期分布"]
    if structure == "dcn":
        pages.append("派息次数分布")
    return pages


def _report_components(structure, samples):
    if structure != "dcn_snowball_combo":
        return {}
    dcn_rows = _component_samples(samples, "dcn")
    snow_rows = _component_samples(samples, "snowball")
    return {
        "dcn": {"summary": _summary(dcn_rows, "dcn"), "charts": _charts(dcn_rows)},
        "snowball": {"summary": _summary(snow_rows, "classic_snowball"), "charts": _charts(snow_rows)},
        "combo": {"summary": _summary(samples, "dcn_snowball_combo"), "charts": _charts(samples)},
    }


def run_backtest(request, cache_path):
    request = validate_request(request)
    frame, source = load_prices(cache_path, request)
    terms = _terms(request)
    structure = request["structure"]
    samples = []
    for position in range(len(frame)):
        if structure == "dcn_snowball_combo":
            dcn = _single(frame, position, _terms(request, request.get("dcn_max_loss")), "dcn")
            snow = _single(frame, position, _terms(request, request.get("snowball_max_loss")), "classic_snowball")
            row = dict(dcn)
            row["dcn_absolute_return"] = dcn.get("absolute_return")
            row["snowball_absolute_return"] = snow.get("absolute_return")
            row["dcn_knocked_in"] = dcn.get("knocked_in")
            row["snowball_knocked_in"] = snow.get("knocked_in")
            if not dcn["still_running"] and not snow["still_running"]:
                row["absolute_return"] = (request["dcn_weight"] * dcn["absolute_return"] +
                                          request["snowball_weight"] * snow["absolute_return"])
                row["annualized_return"] = row["absolute_return"] * 365.0 / max(row["holding_calendar_days"], 1)
            samples.append(row)
        else:
            samples.append(_single(frame, position, terms, structure))
    return {
        "schema_version": 2,
        "report_schema_version": REPORT_SCHEMA_VERSION,
        "engine_version": ENGINE_VERSION,
        "source_commit": SOURCE_COMMIT,
        "source_patch_hash": SOURCE_PATCH_HASH,
        "generated_at": datetime.now().replace(microsecond=0).isoformat(),
        "structure": structure,
        "structure_name": STRUCTURES[structure],
        "index_code": request["index_code"],
        "index_name": INDICES[request["index_code"]],
        "product_name": request["product_name"],
        "report_name": "%s－回测报告" % request["product_name"],
        "time_convention": "calendar_365",
        "day_count_basis": 365,
        "date_adjustment": "following",
        "request": request,
        "source": source,
        "summary": _summary(samples, structure),
        "charts": _charts(samples),
        "report_pages": _report_pages(structure),
        "report_components": _report_components(structure, samples),
        "samples": samples,
    }


def export_excel(result, path):
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "汇总统计"
    summary_sheet.append(["项目", "数值"])
    summary_sheet.append(["产品名称", result.get("product_name")])
    summary_sheet.append(["报告结构版本", result.get("report_schema_version")])
    for key, value in result.get("summary", {}).items():
        summary_sheet.append([key, value])
    summary_sheet.append(["结构", result.get("structure_name")])
    summary_sheet.append(["指数", result.get("index_name")])
    summary_sheet.append(["实际行情起始日", result.get("source", {}).get("actual_start_date")])
    summary_sheet.append(["实际行情截止日", result.get("source", {}).get("actual_end_date")])
    parameter_sheet = workbook.create_sheet("参数与报告信息")
    parameter_sheet.append(["项目", "数值"])
    for key, value in (result.get("request") or {}).items():
        parameter_sheet.append([key, value])
    parameter_sheet.append(["engine_version", result.get("engine_version")])
    parameter_sheet.append(["report_pages", " / ".join(result.get("report_pages") or [])])
    detail = workbook.create_sheet("逐样本明细")
    fields = []
    for row in result.get("samples", []):
        for key in row:
            if key not in fields:
                fields.append(key)
    detail.append(fields)
    for row in result.get("samples", []):
        detail.append([row.get(key) for key in fields])
    os.makedirs(os.path.dirname(path), exist_ok=True)
    workbook.save(path)
