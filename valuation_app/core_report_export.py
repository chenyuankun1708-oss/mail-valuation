import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


MAX_ROWS = 5000


def _safe_rows(value):
    if not isinstance(value, list) or len(value) > MAX_ROWS:
        raise ValueError("核心汇报明细数量无效")
    return value


def _sheet(book, title, headers, rows):
    sheet = book.create_sheet(title)
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="173B6C")
        cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "A2"
    for column in sheet.columns:
        width = min(45, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
        sheet.column_dimensions[column[0].column_letter].width = width
    return sheet


def export_core_report(data):
    if not isinstance(data, dict):
        raise ValueError("核心汇报数据格式无效")
    summary = data.get("summary")
    if not isinstance(summary, dict):
        raise ValueError("缺少核心摘要")
    book = Workbook()
    book.remove(book.active)
    _sheet(book, "核心摘要", ["项目", "数值"], [
        ["核心结论", data.get("narrative")],
        ["开始日", data.get("start")], ["结束日", data.get("end")],
        ["分析范围", data.get("scope")], ["期初资产", summary.get("opening")],
        ["期末资产", summary.get("ending")], ["申购", summary.get("subscriptions")],
        ["赎回", summary.get("redemptions")], ["现金分红", summary.get("dividends")],
        ["区间收益", summary.get("profit")], ["勾稽差额", summary.get("reconciliation")],
        ["可计算数", summary.get("calculated")], ["应计算数", summary.get("eligible")],
        ["覆盖率", summary.get("coverage")], ["说明", data.get("rule_note")],
    ])
    _sheet(book, "重点事项", ["类型", "说明", "触发规则", "严重度"], [
        [row.get("kind"), row.get("text"), row.get("rule"), row.get("level")]
        for row in _safe_rows(data.get("attention", [])) if isinstance(row, dict)
    ])
    _sheet(book, "产品申购赎回", ["日期", "产品", "方向", "金额"], [
        [row.get("date"), row.get("product"), row.get("direction"), row.get("amount")]
        for row in _safe_rows(data.get("product_flows", [])) if isinstance(row, dict)
    ])
    _sheet(book, "收益贡献", ["产品", "收益", "期初资产", "期末资产", "状态"], [
        [row.get("name"), row.get("profit"), (row.get("first") or {}).get("net_assets"),
         (row.get("last") or {}).get("net_assets"), row.get("status")]
        for row in _safe_rows(data.get("contributions", [])) if isinstance(row, dict)
    ])
    concentration_rows = []
    for name, item in (data.get("concentration") or {}).items():
        concentration_rows.append([name, item.get("total"), item.get("top3"), item.get("hhi")])
    _sheet(book, "集中度", ["维度", "总额", "前三大占比", "HHI"], concentration_rows)
    risk = data.get("risk") or {}
    _sheet(book, "风险提示", ["项目", "内容"], [
        ["最大回撤产品", (risk.get("worst_drawdown") or {}).get("name")],
        ["最高波动产品", (risk.get("highest_volatility") or {}).get("name")],
    ])
    _sheet(book, "底层持仓增减估算", ["FOF", "底层产品", "动作", "期初市值", "期末市值", "增减估算", "市值变化"], [
        [row.get("fof"), row.get("name"), row.get("action"), row.get("before"),
         row.get("after"), row.get("position_amount"), row.get("market_change")]
        for row in _safe_rows(data.get("holding_changes", [])) if isinstance(row, dict)
    ])
    _sheet(book, "底层收益率", ["FOF", "底层产品", "策略", "期初日期", "期末日期", "期初价格", "期末价格", "代理收益率", "判断"], [
        [row.get("fof"), row.get("name"), row.get("strategy"), row.get("start_date"),
         row.get("end_date"), row.get("start_price"), row.get("end_price"),
         row.get("return_rate"), row.get("signal")]
        for row in _safe_rows(data.get("holding_returns", [])) if isinstance(row, dict)
    ])
    _sheet(book, "策略收益率", ["一级策略", "期初可计算市值", "持仓数", "代理收益率", "判断"], [
        [row.get("name"), row.get("start_market"), row.get("count"),
         row.get("return_rate"), row.get("signal")]
        for row in _safe_rows(data.get("strategy_returns", [])) if isinstance(row, dict)
    ])
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()
