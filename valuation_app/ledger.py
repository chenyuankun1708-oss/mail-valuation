import math
import os
import re
from dataclasses import asdict, dataclass
from datetime import date, datetime

import openpyxl

from .config import match_product


YEAR_SHEETS = ("2024", "2025", "2026")


@dataclass
class LedgerFlow:
    flow_id: str
    product: str
    flow_date: str
    amount: float
    source_year: str
    source_row: int
    project_text: str
    inferred_date: bool = False
    date_status: str = "recorded"

    def as_dict(self):
        return asdict(self)


def parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        raw = str(int(value))
        if len(raw) == 8 and raw.startswith("20"):
            try:
                return datetime.strptime(raw, "%Y%m%d").date()
            except ValueError:
                return None
    text = str(value or "").strip()
    for fmt in ("%Y.%m.%d", "%Y-%m-%d", "%Y/%m/%d", "%Y%m%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def _amount(value):
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    text = str(value or "").replace(",", "").strip()
    try:
        return float(text)
    except ValueError:
        return None


def _normalize_product(text):
    normalized = str(text or "").replace("F0F", "FOF").replace("f0f", "FOF")
    return match_product(normalized)


def load_ledger(path, snapshots=None):
    """读取年度资金流水。投资金额按万元转为元。"""
    path = os.path.abspath(path)
    values = openpyxl.load_workbook(path, read_only=True, data_only=True)
    formulas = openpyxl.load_workbook(path, read_only=True, data_only=False)
    flows, errors = [], []
    try:
        for sheet_name in YEAR_SHEETS:
            if sheet_name not in values.sheetnames:
                errors.append({"sheet": sheet_name, "error": "缺少年度工作表"})
                continue
            sheet, formula_sheet = values[sheet_name], formulas[sheet_name]
            active_product = None
            for row_index in range(2, sheet.max_row + 1):
                project = sheet.cell(row_index, 3).value
                alias = sheet.cell(row_index, 9).value if sheet.max_column >= 9 else None
                identity = "%s %s" % (project or "", alias or "")
                matched_product = _normalize_product(identity)
                # 年度台账中一个投资事项可能占多行，后续实际投资行会把项目名称
                # 和简称留空。遇到新的非空名称时重置归属；完全空白的续行继承
                # 上一个已匹配事项，不能把这些追加投资漏掉。
                if str(project or "").strip() or str(alias or "").strip():
                    active_product = matched_product
                product = matched_product or (active_product if not str(project or "").strip()
                                              and not str(alias or "").strip() else None)
                amount_value = sheet.cell(row_index, 7).value
                raw_formula = formula_sheet.cell(row_index, 7).value
                amount = _amount(amount_value)
                if not product:
                    continue
                if amount is None:
                    if isinstance(raw_formula, str) and raw_formula.startswith("="):
                        errors.append({"sheet": sheet_name, "row": row_index, "product": product,
                                       "error": "投资金额公式没有缓存结果"})
                    continue
                flow_date = parse_date(sheet.cell(row_index, 6).value)
                flows.append(LedgerFlow(
                    flow_id="%s-%s" % (sheet_name, row_index), product=product,
                    flow_date=flow_date.isoformat() if flow_date else "", amount=round(amount * 10000, 2),
                    source_year=sheet_name, source_row=row_index,
                    project_text=str(project or alias or ""), date_status="recorded" if flow_date else "unresolved",
                ))
    finally:
        values.close()
        formulas.close()
    if snapshots:
        _infer_missing_dates(flows, snapshots)
    return flows, errors


def load_product_metadata(path):
    """读取当前投资本金和2026年度会议审核额度，金额统一为元。"""
    path = os.path.abspath(path)
    values = openpyxl.load_workbook(path, read_only=True, data_only=True)
    metadata, errors = {}, []
    try:
        balance_name = "专户资金余额(2026)"
        if balance_name not in values.sheetnames:
            return {}, [{"sheet": balance_name, "error": "缺少资金余额工作表"}]
        balance = values[balance_name]
        for row in range(3, balance.max_row + 1):
            product = _normalize_product(balance.cell(row, 2).value)
            amount = _amount(balance.cell(row, 4).value)
            if product and amount is not None:
                metadata.setdefault(product, {})["total_investment"] = round(amount * 10000, 2)
        approved = values["2026"]
        for row in range(2, approved.max_row + 1):
            project = approved.cell(row, 3).value
            alias = approved.cell(row, 9).value if approved.max_column >= 9 else None
            product = _normalize_product("%s %s" % (project or "", alias or ""))
            approved_amount = _amount(approved.cell(row, 5).value)
            investment_amount = _amount(approved.cell(row, 7).value)
            investment_date = parse_date(approved.cell(row, 6).value)
            # 排除工作表下方“额度/占用”汇总区。
            if product and approved_amount is not None and (investment_amount is not None or investment_date):
                item = metadata.setdefault(product, {})
                item["approved_quota"] = round(item.get("approved_quota", 0) + approved_amount * 10000, 2)
        for product in metadata:
            metadata[product].setdefault("total_investment", None)
            metadata[product].setdefault("approved_quota", 0)
    finally:
        values.close()
    return metadata, errors


def _infer_missing_dates(flows, snapshots):
    by_product = {}
    for snapshot in snapshots:
        by_product.setdefault(snapshot.product, []).append(snapshot)
    for points in by_product.values():
        points.sort(key=lambda item: item.valuation_date)
    for flow in flows:
        if flow.flow_date:
            continue
        points = by_product.get(flow.product, [])
        candidates = []
        for previous, current in zip(points, points[1:]):
            share_amount = (current.shares - previous.shares) * current.nav
            if not share_amount or math.copysign(1, share_amount) != math.copysign(1, flow.amount):
                continue
            difference = abs(abs(share_amount) - abs(flow.amount))
            if difference <= max(abs(flow.amount) * 0.10, 1000000):
                candidates.append(current.valuation_date)
        if len(candidates) == 1:
            flow.flow_date = candidates[0]
            flow.inferred_date = True
            flow.date_status = "inferred"
        elif points:
            # 无日期流水均为已确认金额；无法匹配份额区间时归到产品最早估值日。
            flow.flow_date = points[0].valuation_date
            flow.inferred_date = True
            flow.date_status = "earliest_valuation"
