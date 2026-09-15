"""Conservative, read-only migration for the historical OTC product workbook."""

import datetime
import hashlib
import json
import os

from openpyxl import load_workbook


def _sha256(path):
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _value(value):
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value.strftime("%Y-%m-%d")
    return value


def _date(value):
    value = _value(value)
    if not value or str(value).strip() in ("--", "存续"):
        return None
    text = str(value).strip().split(" ")[0]
    try:
        datetime.datetime.strptime(text, "%Y-%m-%d")
    except ValueError:
        return None
    return text


def migrate_workbook(store, source, actor="migration"):
    source = os.path.realpath(source)
    if not os.path.isfile(source) or os.path.splitext(source)[1].lower() not in (".xlsx", ".xlsm"):
        raise ValueError("迁移源必须是存在的Excel工作簿")
    before = os.stat(source)
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        rows = []
        for row_number in range(3, 16):
            raw = [_value(sheet.cell(row_number, column).value) for column in range(1, 19)]
            if not raw[0]:
                raise ValueError("第%d行缺少产品名称" % row_number)
            end_date = _date(raw[14])
            status = "存续" if str(raw[14]).strip() == "存续" or str(raw[15]).strip() == "存续" else (
                "已结束原因待补" if end_date else "草稿")
            index_code = str(raw[2] or "").replace(".SH", "").replace(".SZ", "") or None
            reference = {"source_row": row_number, "raw_fields": raw,
                         "recent_dividend_observation": raw[7],
                         "recent_knock_out_observation": raw[10],
                         "reported_end_date": end_date, "reported_end_value": raw[14]}
            terms = {"term_months": int(raw[4]) if raw[4] is not None else None,
                     "initial_level": raw[6], "dividend_level": raw[8],
                     "monthly_dividend": raw[9], "knock_out_level": raw[11],
                     "knock_out_coupon_raw": raw[12], "knock_in_raw": raw[13]}
            rows.append({"source_row": row_number, "status": status, "values": {
                "name": raw[0], "strategy_name": raw[1], "structure": None,
                "index_code": index_code, "notional": float(raw[3]) * 10000 if raw[3] is not None else None,
                "start_date": _date(raw[5]), "terms": terms, "reference": reference,
                "notes": str(raw[17] or "").strip() or None}})
    finally:
        workbook.close()
    after = os.stat(source)
    if (before.st_size, before.st_mtime) != (after.st_size, after.st_mtime):
        raise ValueError("迁移期间源工作簿发生变化")
    return store.import_products(_sha256(source), os.path.basename(source), sheet.title, rows, actor)
