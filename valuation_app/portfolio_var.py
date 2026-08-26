import glob
import math
import os
import re
from collections import defaultdict
from datetime import date, datetime

import openpyxl
import xlrd


EQUITY_FOF_RATE = 0.05
CN_DURATION_BUCKETS = ((1.0, 0.5), (3.0, 2.0), (5.0, 4.0), (float("inf"), 6.0))
RATE_FOF_NAMES = ("国泓资产", "长盛基金汇利2号", "长盛基金汇利海昇5号")
US_TENORS = ((1.0 / 12, "1M"), (0.25, "3M"), (0.5, "6M"), (1, "1Y"),
             (2, "2Y"), (3, "3Y"), (5, "5Y"), (7, "7Y"), (10, "10Y"),
             (20, "20Y"), (30, "30Y"))


def _finite(value):
    try:
        number = float(value)
    except (TypeError, ValueError, OverflowError):
        return None
    return number if math.isfinite(number) else None


def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _excel_date(value, datemode=0):
    if isinstance(value, (datetime, date)):
        return value.date() if isinstance(value, datetime) else value
    number = _finite(value)
    if number and 20000 < number < 80000:
        try:
            return xlrd.xldate_as_datetime(number, datemode).date()
        except (ValueError, TypeError):
            return None
    text = _text(value)
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y%m%d"):
        try:
            return datetime.strptime(text[:10], fmt).date()
        except ValueError:
            pass
    return None


def _quantile(values, level):
    values = sorted(x for x in values if x is not None and math.isfinite(x))
    if not values:
        return None
    position = (len(values) - 1) * level
    low, high = int(math.floor(position)), int(math.ceil(position))
    if low == high:
        return values[low]
    return values[low] + (values[high] - values[low]) * (position - low)


def _yield_var_bp(points, tenor="yield"):
    values = []
    for point in points or []:
        value = _finite(point.get(tenor))
        if value is not None and value > 0:
            values.append(value)
    changes = [(values[index] - values[index - 1]) * 10000.0
               for index in range(1, len(values))]
    return _quantile(changes, 0.99)


def _fx_var(points):
    values = [_finite(point.get("rate")) for point in points or []]
    values = [value for value in values if value is not None and value > 0]
    # A long USD position loses CNY value when USD/CNY falls.
    losses = [max(0.0, -(values[index] / values[index - 1] - 1.0))
              for index in range(1, len(values))]
    return _quantile(losses, 0.99)


def _duration(years):
    years = max(0.0, _finite(years) or 0.0)
    for boundary, duration in CN_DURATION_BUCKETS:
        if years <= boundary:
            return duration
    return 6.0


def _remaining_years(end_date, report_date):
    if not end_date or not report_date:
        return 0.0
    return max(0.0, (end_date - report_date).days / 365.0)


def _credit_bp(rating, layer="", trust=False):
    rating, layer = _text(rating).upper().replace(" ", ""), _text(layer)
    if trust or "次级" in layer or "劣后" in layer:
        return 12.0
    if rating in ("AAA", "AAASF"):
        return 3.0
    if rating in ("AA+", "AA"):
        return 5.0
    if rating in ("A+", "A"):
        return 8.0
    return 10.0


def _nearest_us_tenor(years):
    years = max(1.0 / 12, _finite(years) or 0.0)
    return min(US_TENORS, key=lambda item: abs(item[0] - years))[1]


def _report_date(path, book=None):
    match = re.search(r"(20\d{2})[.\-_年](\d{1,2})[.\-_月](\d{1,2})", os.path.basename(path))
    if match:
        return date(*[int(item) for item in match.groups()])
    return datetime.fromtimestamp(os.path.getmtime(path)).date()


def find_latest_report(env_path=".env"):
    configured = os.environ.get("RISK_REPORT_PATH", "").strip()
    if not configured:
        try:
            with open(env_path, encoding="utf-8") as handle:
                for line in handle:
                    if line.strip().startswith("RISK_REPORT_PATH="):
                        configured = line.split("=", 1)[1].strip().strip("'\"")
                        break
        except OSError:
            pass
    if configured and os.path.isfile(configured):
        return configured
    roots = [os.path.abspath("data_sources"), os.path.join(os.path.expanduser("~"), "Downloads")]
    matches = []
    for root in roots:
        for extension in ("xls", "xlsx"):
            matches.extend(glob.glob(os.path.join(root, "风控日报*." + extension)))
    return max(matches, key=os.path.getmtime) if matches else None


def _asset(identifier, product, name, category, amount, amount_basis, currency="CNY", **extra):
    item = {"id": identifier, "product": product or name, "name": name,
            "category": category, "currency": currency, "exposure": amount,
            "source_amount": amount, "source_currency": currency,
            "amount_basis": amount_basis, "status": "pending", "error": None}
    item.update(extra)
    return item


def parse_risk_report(path):
    book = xlrd.open_workbook(path, on_demand=True)
    report_date = _report_date(path, book)
    assets = []
    names = set(book.sheet_names())
    if "基金专户一览表" in names:
        sheet = book.sheet_by_name("基金专户一览表")
        for row in range(1, sheet.nrows):
            name = _text(sheet.cell_value(row, 1))
            amount = _finite(sheet.cell_value(row, 3))
            if not name or amount is None or any(key in name for key in RATE_FOF_NAMES):
                continue
            assets.append(_asset("FOF-%d" % row, name, name, "权益FOF", amount * 10000.0,
                                 "日报产品本金", fixed_var_rate=EQUITY_FOF_RATE))
    detail_sheets = [name for name in names if name.startswith("基金专户一览表（")]
    for sheet_index, sheet_name in enumerate(detail_sheets):
        sheet = book.sheet_by_name(sheet_name)
        header = [_text(sheet.cell_value(0, col)) for col in range(sheet.ncols)]
        amount_col = next((i for i, value in enumerate(header) if "投资金额" in value), None)
        end_col = next((i for i, value in enumerate(header) if "结束日期" in value or "到期" in value), None)
        rating_col = next((i for i, value in enumerate(header) if "评级" in value), None)
        code_col = next((i for i, value in enumerate(header) if "债券代码" in value), None)
        name_col = next((i for i, value in enumerate(header) if "债券简称" in value), None)
        product = sheet_name.split("（", 1)[1].rstrip("）")
        if amount_col is None:
            continue
        for row in range(1, sheet.nrows):
            raw_amount = sheet.cell_value(row, amount_col)
            text_amount = _text(raw_amount).upper()
            amount = _finite(raw_amount)
            currency = "CNY"
            if text_amount.startswith("USD"):
                amount = _finite(re.sub(r"[^0-9.\-]", "", text_amount))
                amount = amount * 1000000.0 if amount is not None else None
                currency = "USD"
            elif amount is not None:
                amount *= 10000.0
            if amount is None or amount <= 0:
                continue
            bond_name = _text(sheet.cell_value(row, name_col)) if name_col is not None else ""
            code = _text(sheet.cell_value(row, code_col)) if code_col is not None else ""
            end = _excel_date(sheet.cell_value(row, end_col), book.datemode) if end_col is not None else None
            assets.append(_asset("RATE-%s-%d" % (sheet_index, row), product,
                                 bond_name or code or (product + "底层资产"), "利率专户底层",
                                 amount, "日报底层投资金额", currency=currency,
                                 code=code, rating=_text(sheet.cell_value(row, rating_col)) if rating_col is not None else "",
                                 layer="", end_date=end.isoformat() if end else None))
    if "结构化产品一览表" in names:
        sheet = book.sheet_by_name("结构化产品一览表")
        for row in range(2, sheet.nrows):
            amount = _finite(sheet.cell_value(row, 6))
            name = _text(sheet.cell_value(row, 3))
            if not name or amount is None or amount <= 0:
                continue
            end = _excel_date(sheet.cell_value(row, 14), book.datemode)
            assets.append(_asset("STRUCT-%d" % row, name, name, "结构化固收", amount,
                                 "日报认购金额", rating=_text(sheet.cell_value(row, 5)),
                                 layer=_text(sheet.cell_value(row, 4)),
                                 end_date=end.isoformat() if end else None))
    if "信托产品一览表" in names:
        sheet = book.sheet_by_name("信托产品一览表")
        for row in range(2, sheet.nrows):
            amount = _finite(sheet.cell_value(row, 5))
            name = _text(sheet.cell_value(row, 1))
            if not name or amount is None or amount <= 0:
                continue
            end = _excel_date(sheet.cell_value(row, 10), book.datemode)
            assets.append(_asset("TRUST-%d" % row, name, name, "信托固收", amount * 10000.0,
                                 "日报认购金额", rating="NR", layer="信托",
                                 end_date=end.isoformat() if end else None, trust=True))
    if "挂钩结构化资产标的收益互换" in names:
        sheet = book.sheet_by_name("挂钩结构化资产标的收益互换")
        for row in range(2, sheet.nrows):
            amount = _finite(sheet.cell_value(row, 8))
            name, code = _text(sheet.cell_value(row, 4)), _text(sheet.cell_value(row, 3))
            if amount is None or amount <= 0 or not (name or code):
                continue
            end = _excel_date(sheet.cell_value(row, 7), book.datemode)
            issuer = _text(sheet.cell_value(row, 5))
            is_us = code.startswith("US") or code.startswith("XS") or "GNR " in name or "FNR " in name
            agency = "GNR " in name or "FNR " in name or "Ginnie" in issuer or "Fannie" in issuer
            assets.append(_asset("TRS-%d" % row, "挂钩结构化资产标的收益互换", name or code,
                                 "美元Agency MBS" if agency else ("美元海外ABS" if is_us else "境内收益互换固收"),
                                 amount, "日报期末人民币名义本金", currency="USD" if is_us else "CNY",
                                 code=code, rating="", layer="", agency_mbs=agency,
                                 end_date=end.isoformat() if end else None, amount_is_cny=True,
                                 source_currency="CNY"))
    if "碳金融业务" in names:
        sheet = book.sheet_by_name("碳金融业务")
        for row in range(1, sheet.nrows):
            name, market = _text(sheet.cell_value(row, 2)), _text(sheet.cell_value(row, 3))
            quantity = _finite(sheet.cell_value(row, 5))
            if name and quantity is not None and quantity > 0:
                assets.append(_asset("CARBON-%d" % row, "碳金融业务", name, "碳资产", None,
                                     "持仓数量×碳价文件最新价格", quantity=quantity, market=market))
    book.release_resources()
    return {"report_date": report_date.isoformat(), "report_file": os.path.basename(path), "assets": assets}


def _workbook_rows(path):
    if path.lower().endswith(".xls"):
        book = xlrd.open_workbook(path, on_demand=True)
        try:
            return [(sheet.name, [[sheet.cell_value(row, col) for col in range(sheet.ncols)]
                                  for row in range(sheet.nrows)], book.datemode) for sheet in book.sheets()]
        finally:
            book.release_resources()
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return [(sheet.title, [list(row) for row in sheet.iter_rows(values_only=True)], 0)
                for sheet in book.worksheets]
    finally:
        book.close()


def load_carbon_prices(root="data_sources"):
    paths = []
    roots = [root]
    if os.path.normcase(os.path.normpath(root)) == os.path.normcase(os.path.normpath("data_sources")):
        roots.append(".")
    for search_root in roots:
        for extension in ("xls", "xlsx"):
            paths.extend(glob.glob(os.path.join(search_root, "碳排放价格*." + extension)))
    if not paths:
        raise IOError("未找到data_sources或项目根目录下的碳排放价格.xls[x]")
    path = max(paths, key=os.path.getmtime)
    series = defaultdict(dict)
    aliases = {"date": ("日期", "交易日", "时间"), "market": ("市场", "品种", "标的", "名称"),
               "price": ("收盘价", "价格", "成交价", "成交均价", "均价")}
    for sheet_name, rows, datemode in _workbook_rows(path):
        header_index, columns = None, {}
        for index, row in enumerate(rows[:20]):
            texts = [_text(value) for value in row]
            found = {key: next((i for i, value in enumerate(texts)
                                if any(alias in value for alias in values)), None)
                     for key, values in aliases.items()}
            if found["date"] is not None and found["price"] is not None:
                header_index, columns = index, found
                break
        if header_index is None:
            continue
        for row in rows[header_index + 1:]:
            if max(columns["date"], columns["price"]) >= len(row):
                continue
            day = _excel_date(row[columns["date"]], datemode)
            price = _finite(row[columns["price"]])
            market = (_text(row[columns["market"]]) if columns["market"] is not None and
                      columns["market"] < len(row) else sheet_name)
            if day and price is not None and price > 0:
                key = "CCER" if "CCER" in market.upper() or "核证" in market else ("重庆" if "重庆" in market else market)
                series[key][day.isoformat()] = price
    normalized = {key: [{"date": day, "price": values[day]} for day in sorted(values)]
                  for key, values in series.items() if values}
    if not normalized:
        raise ValueError("碳排放价格文件没有可识别的日期和价格序列")
    return {"file": os.path.basename(path), "series": normalized}


def _carbon_series(carbon, asset):
    key = "CCER" if "CCER" in asset["name"].upper() or "核证" in asset["name"] else ("重庆" if "重庆" in asset.get("market", "") or "重庆" in asset["name"] else None)
    series = carbon.get("series", {})
    matched = (series.get(key) or []) if key else []
    if matched:
        return matched
    # The approved management proxy is the return series in the local workbook,
    # irrespective of the carbon instrument's own market or variety.
    if len(series) == 1:
        return next(iter(series.values()))
    return []


def _summaries(assets):
    by_product = defaultdict(lambda: {"exposure": 0.0, "var": 0.0, "assets": 0})
    by_category = defaultdict(lambda: {"exposure": 0.0, "var": 0.0, "assets": 0})
    calculated_exposure = total_exposure = total_var = 0.0
    calculated_assets = 0
    for item in assets:
        exposure = item.get("exposure") or 0.0
        total_exposure += exposure
        for target, key in ((by_product, item["product"]), (by_category, item["category"])):
            target[key]["exposure"] += exposure
            target[key]["var"] += item.get("total_var") or 0.0
            target[key]["assets"] += 1
        if item["status"] == "calculated":
            calculated_assets += 1
            calculated_exposure += exposure
            total_var += item.get("total_var") or 0.0
    def rows(values, key_name):
        return [dict({key_name: key}, **values[key]) for key in sorted(values)]
    return {"product_summary": rows(by_product, "product"),
            "category_summary": rows(by_category, "category"),
            "summary": {"assets": len(assets), "total_exposure": total_exposure,
                        "calculated_assets": calculated_assets,
                        "missing_assets": len(assets) - calculated_assets,
                        "asset_coverage_pct": calculated_assets / len(assets) * 100.0 if assets else 0.0,
                        "calculated_exposure": calculated_exposure,
                        "missing_exposure": total_exposure - calculated_exposure,
                        "coverage_pct": calculated_exposure / total_exposure * 100.0 if total_exposure else 0.0,
                        "total_var": total_var}}


def calculate(report, risk_cache, carbon=None):
    report_date = date.fromisoformat(report["report_date"])
    cn_points = risk_cache.get("yield_curves", {}).get("CN10Y", {}).get("points", [])
    cn_bp = _yield_var_bp(cn_points)
    us_points = risk_cache.get("yield_curves", {}).get("US_TREASURY", {}).get("points", [])
    fx_points = risk_cache.get("exchange_rates", {}).get("USDCNY", {}).get("points", [])
    fx_rate = _finite(fx_points[-1].get("rate")) if fx_points else None
    fx_var = _fx_var(fx_points)
    assets = []
    for original in report["assets"]:
        item = dict(original)
        item.update({"risk_free_var": None, "credit_var": None, "fx_var": None,
                     "carbon_var": None, "total_var": None})
        if item.get("fixed_var_rate") is not None:
            item["total_var"] = item["exposure"] * item["fixed_var_rate"]
            item["status"] = "calculated"
            item["method"] = "权益FOF固定5%"
            item["calculation_formula"] = "敞口代理×固定比例5%"
        elif item["category"] == "碳资产":
            points = _carbon_series(carbon or {}, item)
            prices = [_finite(point.get("price")) for point in points]
            prices = [value for value in prices if value is not None and value > 0]
            losses = [max(0.0, -(prices[index] / prices[index - 1] - 1.0))
                      for index in range(1, len(prices))]
            rate = _quantile(losses, 0.99)
            if prices and rate is not None:
                item["latest_price"] = prices[-1]
                item["exposure"] = item["quantity"] * prices[-1]
                item["carbon_var"] = item["exposure"] * rate
                item["carbon_var_rate"] = rate
                item["total_var"] = item["carbon_var"]
                item["status"] = "calculated"
                item["method"] = "本地Excel统一碳价代理左尾99%"
                item["calculation_formula"] = "持仓数量×Excel最新碳价×统一碳价代理99%历史日跌幅"
            else:
                item["status"], item["error"] = "missing", "没有唯一匹配的碳价历史序列"
                item["method"] = "碳价历史跌幅左尾99%"
                item["calculation_formula"] = "持仓数量×最新碳价×99%历史日跌幅"
        else:
            end = date.fromisoformat(item["end_date"]) if item.get("end_date") else report_date
            years = _remaining_years(end, report_date)
            is_us = item["currency"] == "USD"
            if item.get("agency_mbs"):
                duration, credit_bp = 4.0, 5.0
            elif is_us:
                duration, credit_bp = 3.0, 10.0
            else:
                duration = _duration(years)
                credit_bp = _credit_bp(item.get("rating"), item.get("layer"), item.get("trust", False))
            exposure = item.get("exposure")
            if is_us and not item.get("amount_is_cny"):
                exposure = exposure * fx_rate if exposure is not None and fx_rate is not None else None
            tenor = "10Y"
            rate_bp = _yield_var_bp(us_points, tenor) if is_us else cn_bp
            item.update({"exposure": exposure, "remaining_years": round(years, 4),
                         "proxy_duration": duration, "credit_spread_bp": credit_bp,
                         "rate_tenor": tenor, "rate_var_bp": rate_bp,
                         "fx_spot": fx_rate if is_us else None,
                         "fx_var_rate": fx_var if is_us else None})
            if exposure is None or rate_bp is None:
                item["status"], item["error"] = "missing", "缺少金额、汇率或利率曲线"
            else:
                item["risk_free_var"] = exposure * duration * rate_bp / 10000.0
                item["credit_var"] = exposure * duration * credit_bp / 10000.0
                if is_us and fx_var is not None:
                    item["fx_var"] = exposure * fx_var
                item["total_var"] = sum(value or 0.0 for value in
                                        (item["risk_free_var"], item["credit_var"], item["fx_var"]))
                item["status"] = "calculated"
            item["method"] = "美元利率+固定信用利差+汇率" if is_us else "中债10Y+固定信用利差"
            item["calculation_formula"] = ("敞口代理×代理久期×(美国国债10Y VaR bp+信用冲击bp)/10000+敞口代理×USD/CNY VaR"
                                           if is_us else
                                           "敞口代理×代理久期×(中债10Y VaR bp+信用冲击bp)/10000")
        assets.append(item)
    summaries = _summaries(assets)
    carbon_series = []
    if carbon and len(carbon.get("series", {})) == 1:
        carbon_series = next(iter(carbon["series"].values()))
    return {"report_date": report["report_date"], "report_file": report["report_file"],
            "assets": assets, "product_summary": summaries["product_summary"],
            "category_summary": summaries["category_summary"],
            "summary": summaries["summary"],
            "parameters": {"equity_fof_rate": EQUITY_FOF_RATE, "cn10y_var_bp": cn_bp,
                           "usd_cny_var_rate": fx_var, "carbon_price_points": carbon_series,
                           "carbon_proxy_scope": "本地Excel单一涨跌幅序列统一代理全部碳资产"},
            "status": "ok", "error": None}


def export_assets_excel(portfolio, path):
    """Export the auditable asset-level VaR result used by the web page."""
    import openpyxl
    headers = [
        ("类别", "category"), ("产品", "product"), ("资产", "name"), ("代码", "code"),
        ("币种", "currency"), ("原始金额", "source_amount"), ("原始币种", "source_currency"),
        ("金额口径", "amount_basis"), ("人民币敞口代理", "exposure"), ("评级", "rating"),
        ("层级", "layer"), ("到期日", "end_date"), ("剩余期限（年）", "remaining_years"),
        ("代理久期", "proxy_duration"), ("利率基准", "rate_tenor"),
        ("利率VaR（bp）", "rate_var_bp"), ("利率VaR金额", "risk_free_var"),
        ("信用冲击（bp）", "credit_spread_bp"), ("信用VaR金额", "credit_var"),
        ("USD/CNY", "fx_spot"), ("汇率VaR率", "fx_var_rate"), ("汇率VaR金额", "fx_var"),
        ("碳持仓数量", "quantity"), ("最新碳价", "latest_price"),
        ("碳价VaR率", "carbon_var_rate"), ("碳价VaR金额", "carbon_var"),
        ("固定比例", "fixed_var_rate"), ("计算公式", "calculation_formula"),
        ("合计VaR", "total_var"), ("状态", "status"), ("错误", "error"), ("方法", "method"),
    ]
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "全部资产VaR明细"
    sheet.append([label for label, _ in headers])
    for asset in portfolio.get("assets", []):
        sheet.append([asset.get(key) for _, key in headers])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = openpyxl.styles.Font(bold=True)
    for column in sheet.columns:
        letter = column[0].column_letter
        sheet.column_dimensions[letter].width = min(42, max(10, max(len(str(c.value or "")) for c in column) + 2))
    absolute = os.path.abspath(path)
    os.makedirs(os.path.dirname(absolute), exist_ok=True)
    book.save(absolute)
    book.close()
    return absolute


def update_portfolio_var(previous, risk_cache, attempted, env_path=".env"):
    old = previous.get("portfolio_var", {})
    try:
        path = find_latest_report(env_path)
        if not path:
            raise IOError("未找到风控日报*.xls[x]")
        report = parse_risk_report(path)
        carbon, carbon_error = None, None
        try:
            carbon = load_carbon_prices()
        except Exception as exc:
            carbon_error = "%s: %s" % (type(exc).__name__, exc)
        result = calculate(report, risk_cache, carbon)
        if carbon_error and old.get("assets"):
            old_carbon = {item.get("name"): item for item in old["assets"]
                          if item.get("category") == "碳资产" and item.get("status") == "calculated"}
            for index, item in enumerate(result["assets"]):
                if item.get("category") == "碳资产" and item.get("status") != "calculated" and item.get("name") in old_carbon:
                    cached = dict(old_carbon[item["name"]])
                    cached["used_cache"] = True
                    cached["error"] = carbon_error
                    result["assets"][index] = cached
            summaries = _summaries(result["assets"])
            result.update(summaries)
        result["updated_at"] = attempted
        result["carbon_source"] = carbon.get("file") if carbon else None
        result["carbon_error"] = carbon_error
        return result, None
    except Exception as exc:
        result = dict(old) if old.get("assets") else {"assets": [], "product_summary": [],
            "category_summary": [], "summary": {}, "status": "missing"}
        result["error"] = "%s: %s" % (type(exc).__name__, exc)
        result["used_cache"] = bool(old.get("assets"))
        return result, {"error": result["error"], "used_cache": result["used_cache"]}
