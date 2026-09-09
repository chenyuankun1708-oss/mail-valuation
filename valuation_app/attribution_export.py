import io
import math

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


MAX_DETAIL_ROWS = 20000
MAX_SUMMARY_ROWS = 10000


def _rows(value, limit):
    if not isinstance(value, list) or len(value) > limit:
        raise ValueError("归因明细数量无效")
    return [row for row in value if isinstance(row, dict)]


def _cell(value):
    if isinstance(value, float) and not math.isfinite(value):
        return None
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        return "'" + value
    return value


def _sheet(book, title, headers, rows):
    sheet = book.create_sheet(title)
    sheet.append(headers)
    for row in rows:
        sheet.append([_cell(value) for value in row])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="173B6C")
        cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "A2"
    for column in sheet.columns:
        width = min(45, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
        sheet.column_dimensions[column[0].column_letter].width = width


def _summary_rows(rows):
    return [[row.get("name"), row.get("price_impact"), row.get("position_amount"),
             row.get("market_change"), row.get("formula_difference"), row.get("count")]
            for row in rows]


def export_attribution(data):
    if not isinstance(data, dict) or not isinstance(data.get("summary"), dict):
        raise ValueError("归因数据格式无效")
    detail = _rows(data.get("intervals", []), MAX_DETAIL_ROWS)
    products = _rows(data.get("products", []), MAX_SUMMARY_ROWS)
    underlying = _rows(data.get("underlying", []), MAX_SUMMARY_ROWS)
    strategies = _rows(data.get("strategies", []), MAX_SUMMARY_ROWS)
    managers = _rows(data.get("managers", []), MAX_SUMMARY_ROWS)
    actions = _rows(data.get("actions", []), MAX_SUMMARY_ROWS)
    summary = data["summary"]
    book = Workbook()
    book.remove(book.active)
    _sheet(book, "归因概览", ["项目", "数值"], [
        ["页面名称", "持仓变动归因（估算）"], ["开始日", data.get("start")],
        ["结束日", data.get("end")], ["分析范围", data.get("scope")],
        ["价格影响", summary.get("price_impact")], ["仓位变化金额", summary.get("position_amount")],
        ["底层市值变化", summary.get("market_change")], ["公式勾稽差额", summary.get("formula_difference")],
        ["顶层产品区间收益", summary.get("product_profit")],
        ["现金费用及其他未归属", summary.get("unattributed")],
        ["可归属", summary.get("attributable")], ["部分可归属", summary.get("partial")],
        ["不可归属", summary.get("unavailable")],
        ["口径说明", data.get("rule_note")],
    ])
    _sheet(book, "产品勾稽", ["FOF", "区间收益", "价格影响", "现金费用及其他未归属", "期初资产", "期末资产"], [
        [row.get("name"), row.get("profit"), row.get("price_impact"), row.get("unattributed"),
         row.get("opening"), row.get("ending")] for row in products
    ])
    headers = ["名称", "价格影响", "仓位变化金额", "市值变化", "公式差额", "明细数"]
    _sheet(book, "底层汇总", headers, _summary_rows(underlying))
    _sheet(book, "策略汇总", headers, _summary_rows(strategies))
    _sheet(book, "管理人汇总", headers, _summary_rows(managers))
    _sheet(book, "动作汇总", headers, _summary_rows(actions))
    _sheet(book, "逐区间明细", ["FOF", "底层产品", "代码", "一级策略", "管理人", "开始日", "结束日", "动作",
                                "期初份额q0", "期初价格P0", "期末份额q1", "期末价格P1", "期初市值", "期末市值",
                                "价格影响q0×(P1-P0)", "仓位变化(q1-q0)×P1", "市值变化", "公式差额", "归属状态", "说明"], [
        [row.get("fof"), row.get("name"), row.get("code"), row.get("strategy"), row.get("manager"),
         row.get("start_date"), row.get("end_date"), row.get("action"), row.get("q0"), row.get("p0"),
         row.get("q1"), row.get("p1"), row.get("mv0"), row.get("mv1"), row.get("price_impact"),
         row.get("position_amount"), row.get("market_change"), row.get("formula_difference"),
         row.get("quality"), row.get("note")] for row in detail
    ])
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()
